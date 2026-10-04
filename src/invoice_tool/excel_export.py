from __future__ import annotations

import os
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import InvoiceFolder, InvoiceRecord
from .validation import validate_application, validate_invoice

HEADERS = [
    "申请单编号", "申请人", "申请日期", "名称", "型号规格", "单位", "数量",
    "单价", "金额", "税额", "价税合计", "发票号码", "付款日期",
]
WIDTHS = [11.8888888888889, 8.88888888888889, 12, 18, 18, 8, 8, 10, 10, 10, 11, 22, 12]
GRAY_FILL = "BFBFBF"


def default_output_name() -> str:
    return f"采购记录_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def export_excel(folders: list[InvoiceFolder], output_path: str | Path) -> Path:
    """Export confirmed application/invoice groups to the 13-column purchase ledger."""
    # Empty hand-created folders are an organising aid, not a purchase record.
    # They must not block a valid application from being exported.
    folders = [folder for folder in folders if folder.entries]
    if not folders:
        raise ValueError("没有可导出的采购单。")
    numbers = [folder.application.application_number.strip() for folder in folders if folder.application]
    duplicates = {number for number in numbers if number and numbers.count(number) > 1}
    for folder in folders:
        records = _records(folder)
        if not records:
            raise ValueError("采购单中没有已识别的发票。")
        result = validate_application(
            folder.application, records,
            duplicate_number=bool(folder.application and folder.application.application_number.strip() in duplicates),
        )
        if not result.valid:
            raise ValueError("；".join(issue.message for issue in result.errors))
        if not folder.application or not folder.application.confirmed:
            raise ValueError("请先确认每份申请表信息。")
        for invoice in records:
            result = validate_invoice(invoice)
            if not result.valid:
                raise ValueError("存在未解决的发票校验问题，不能导出。")
            if not invoice.confirmed:
                raise ValueError("所有发票必须先人工确认。")

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    _style_template(sheet)
    row = 2
    for folder in folders:
        application = folder.application
        assert application is not None
        application_start = row
        for invoice in _records(folder):
            invoice_start = row
            for item in invoice.items:
                _write_item_row(sheet, row, application.application_number, application.applicant, application.purchase_date,
                                application.payment_date, invoice, item)
                row += 1
            if row - 1 > invoice_start:
                _merge_and_center(sheet, invoice_start, row - 1, (11, 12))
        if row - 1 > application_start:
            _merge_and_center(sheet, application_start, row - 1, (1, 2, 3, 13))
    _set_formats(sheet, row - 1)
    with tempfile.NamedTemporaryFile(prefix="purchase-export-", suffix=".xlsx", dir=target.parent, delete=False) as temporary:
        temp_path = Path(temporary.name)
    try:
        workbook.save(temp_path)
        _verify_export(temp_path, folders)
        os.replace(temp_path, target)
    finally:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
    return target


def _records(folder: InvoiceFolder) -> list[InvoiceRecord]:
    return [entry.record for entry in folder.entries if entry.record is not None]


def _style_template(sheet) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name="宋体", size=11, bold=True)
    data_font = Font(name="宋体", size=11)
    gray = PatternFill("solid", fgColor=GRAY_FILL)
    for column, (header, width) in enumerate(zip(HEADERS, WIDTHS, strict=True), start=1):
        cell = sheet.cell(1, column, header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        if column in (8, 11):
            cell.fill = gray
        sheet.column_dimensions[cell.column_letter].width = width


def _write_item_row(sheet, row: int, application_number: str, applicant: str, purchase_date: date | None,
                    payment_date: date | None, invoice: InvoiceRecord, item) -> None:
    _style_data_row(sheet, row)
    values = [application_number, applicant, purchase_date, item.project_name, item.specification, item.unit,
              _number(item.quantity), _number(item.unit_price), _number(item.amount), _number(item.tax_amount),
              _number(invoice.total_with_tax), invoice.invoice_number, payment_date]
    for column, value in enumerate(values, start=1):
        sheet.cell(row, column, value)
    for column in (1, 12):
        cell = sheet.cell(row, column)
        cell.quotePrefix = True
        cell.number_format = "@"


def _style_data_row(sheet, row: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    gray = PatternFill("solid", fgColor=GRAY_FILL)
    for column in range(1, 14):
        cell = sheet.cell(row, column)
        cell.font = Font(name="宋体", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        if column in (8, 11):
            cell.fill = gray


def _merge_and_center(sheet, start: int, end: int, columns: tuple[int, ...]) -> None:
    for column in columns:
        sheet.merge_cells(start_row=start, start_column=column, end_row=end, end_column=column)
        sheet.cell(start, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_formats(sheet, final_row: int) -> None:
    for row in range(2, final_row + 1):
        sheet.cell(row, 3).number_format = "yyyy-mm-dd"
        sheet.cell(row, 13).number_format = "yyyy-mm-dd"
        quantity = sheet.cell(row, 7).value
        sheet.cell(row, 7).number_format = "0" if quantity is not None and float(quantity).is_integer() else "0.##############"
        sheet.cell(row, 8).number_format = "0.##############"
        for column in (9, 10, 11):
            sheet.cell(row, column).number_format = "0.00"
        for column in (1, 12):
            sheet.cell(row, column).number_format = "@"


def _number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _verify_export(path: Path, folders: list[InvoiceFolder]) -> None:
    workbook = load_workbook(path, data_only=False)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError("导出的工作表结构错误。")
    sheet = workbook["Sheet1"]
    expected_rows = 1 + sum(len(invoice.items) for folder in folders for invoice in _records(folder))
    if sheet.max_row != expected_rows or sheet.max_column != 13:
        raise ValueError("导出的明细行数或列数错误。")
    if [cell.value for cell in sheet[1]] != HEADERS:
        raise ValueError("导出的表头错误。")
    row = 2
    for folder in folders:
        application = folder.application
        assert application is not None
        application_start = row
        for invoice in _records(folder):
            invoice_start, invoice_end = row, row + len(invoice.items) - 1
            if str(sheet.cell(invoice_start, 12).value) != invoice.invoice_number:
                raise ValueError("导出的发票号码错误。")
            if sheet.cell(invoice_start, 12).data_type != "s" or not sheet.cell(invoice_start, 12).quotePrefix:
                raise ValueError("发票号码没有按强制文本格式导出。")
            if abs(Decimal(str(sheet.cell(invoice_start, 11).value)) - invoice.total_with_tax) > Decimal("0.0001"):
                raise ValueError("导出的价税合计错误。")
            if invoice_end > invoice_start:
                merged = {str(value) for value in sheet.merged_cells.ranges}
                if not {f"K{invoice_start}:K{invoice_end}", f"L{invoice_start}:L{invoice_end}"}.issubset(merged):
                    raise ValueError("导出的发票合并单元格错误。")
            row = invoice_end + 1
        if str(sheet.cell(application_start, 1).value) != application.application_number:
            raise ValueError("导出的申请单编号错误。")
        if sheet.cell(application_start, 1).data_type != "s" or not sheet.cell(application_start, 1).quotePrefix:
            raise ValueError("申请单编号没有按强制文本格式导出。")
        exported_date = sheet.cell(application_start, 3).value
        if isinstance(exported_date, datetime):
            exported_date = exported_date.date()
        if exported_date != application.purchase_date:
            raise ValueError("导出的申请日期错误。")
        if row - 1 > application_start:
            merged = {str(value) for value in sheet.merged_cells.ranges}
            expected = {f"{column}{application_start}:{column}{row - 1}" for column in "ABCM"}
            if not expected.issubset(merged):
                raise ValueError("导出的申请单合并单元格错误。")
