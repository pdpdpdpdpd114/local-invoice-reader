from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from invoice_tool.excel_export import HEADERS, export_excel
from invoice_tool.models import InvoiceEntry, InvoiceFolder, InvoiceItem, InvoiceRecord, PurchaseApplication


def _invoice(number: str, total: str, items: list[InvoiceItem]) -> InvoiceRecord:
    amount = sum((item.amount or Decimal("0")) for item in items)
    tax = sum((item.tax_amount or Decimal("0")) for item in items)
    return InvoiceRecord(number, Decimal(total), items=items, amount_total=amount, tax_total=tax, confirmed=True)


def _folder(application_number: str, invoices: list[InvoiceRecord]) -> InvoiceFolder:
    application = PurchaseApplication(
        applicant="杨涛义",
        purchase_date=date(2026, 7, 29),
        group_name="车架",
        purchase_amount=sum((invoice.total_with_tax or Decimal("0")) for invoice in invoices),
        application_number=application_number,
        confirmed=True,
    )
    return InvoiceFolder(
        purchase_date=application.purchase_date,
        group_name=application.group_name,
        application=application,
        entries=[InvoiceEntry(source_file=Path(f"{invoice.invoice_number}.pdf"), record=invoice) for invoice in invoices],
    )


def test_export_matches_purchase_ledger_headers_and_merges(tmp_path: Path) -> None:
    invoice = _invoice(
        "00000000000000000022",
        "22.22",
        [
            InvoiceItem("项目A" * 20, "", "件", Decimal("1"), Decimal("10"), Decimal("10"), Decimal("0.01"), Decimal("0.10")),
            InvoiceItem("项目B", "规格", "件", Decimal("2"), Decimal("6"), Decimal("12"), Decimal("0.01"), Decimal("0.12")),
        ],
    )
    output = export_excel([_folder("CJ20260729", [invoice])], tmp_path / "purchase.xlsx")
    sheet = load_workbook(output)["Sheet1"]

    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_column == 13
    assert sheet["A2"].value == "CJ20260729"
    assert sheet["A2"].data_type == "s" and sheet["A2"].quotePrefix is True
    assert sheet["B2"].value == "杨涛义"
    assert sheet["C2"].value.date() == date(2026, 7, 29)
    assert sheet["K2"].value == 22.22
    assert sheet["L2"].value == "00000000000000000022"
    assert sheet["L2"].data_type == "s" and sheet["L2"].quotePrefix is True
    assert {"A2:A3", "B2:B3", "C2:C3", "M2:M3", "K2:K3", "L2:L3"}.issubset(
        {str(item) for item in sheet.merged_cells.ranges}
    )
    assert sheet["H2"].fill.fill_type == "solid"
    assert sheet["K2"].fill.fill_type == "solid"
    assert sheet["M2"].value is None


def test_export_keeps_application_then_invoice_order(tmp_path: Path) -> None:
    first = _invoice(
        "00000000000000000022", "10.10",
        [InvoiceItem("第一张", "", "件", Decimal("1"), Decimal("10"), Decimal("10"), Decimal("0.01"), Decimal("0.10"))],
    )
    second = _invoice(
        "00000000000000000023", "20.20",
        [InvoiceItem("第二张", "", "件", Decimal("2"), Decimal("10"), Decimal("20"), Decimal("0.01"), Decimal("0.20"))],
    )
    output = export_excel([_folder("CJ20260729", [first]), _folder("CJ20260730", [second])], tmp_path / "ordered.xlsx")
    sheet = load_workbook(output)["Sheet1"]
    assert [sheet.cell(row, 1).value for row in (2, 3)] == ["CJ20260729", "CJ20260730"]
    assert [sheet.cell(row, 12).value for row in (2, 3)] == [first.invoice_number, second.invoice_number]
